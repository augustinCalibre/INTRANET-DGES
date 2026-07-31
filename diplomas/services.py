import csv
import io
import unicodedata
from pathlib import Path

from django.contrib.auth.models import User
from django.urls import reverse

from accounts.constants import ROLE_ADMINISTRATEUR, ROLE_DIRECTEUR_GENERAL
from core.models import Notification

from .models import Diplome, DiplomeHistory, LotHistory

MAX_IMPORT_ROWS = 2000
SUPPORTED_IMPORT_EXTENSIONS = (".csv", ".xlsx")

COLUMN_ALIASES = {
    "nom_beneficiaire": (
        "nom",
        "noms",
        "nom beneficiaire",
        "nom du beneficiaire",
        "beneficiaire",
        "nom et prenom",
        "nom et prenoms",
        "nom complet",
    ),
    "numero_diplome": (
        "numero",
        "numero diplome",
        "numero du diplome",
        "numero de diplome",
        "no diplome",
        "n diplome",
        "reference",
        "reference diplome",
    ),
    "filiere": ("filiere", "option", "section", "domaine", "faculte"),
    "etablissement": ("etablissement", "universite", "ecole", "institution", "provenance"),
    "annee_academique": (
        "annee",
        "annee academique",
        "annee scolaire",
        "promotion",
        "exercice",
    ),
    "observations": ("observation", "observations", "remarque", "remarques", "commentaire"),
}

REQUIRED_COLUMNS = ("nom_beneficiaire", "numero_diplome")

FIELD_MAX_LENGTHS = {
    "nom_beneficiaire": 180,
    "numero_diplome": 80,
    "filiere": 150,
    "etablissement": 180,
    "annee_academique": 20,
}


# ---------------------------------------------------------------- historique


def register_lot_history(lot, user, action, previous_status="", current_status="", commentaire=""):
    return LotHistory.objects.create(
        lot=lot,
        action=action,
        ancien_statut=previous_status,
        nouveau_statut=current_status,
        commentaire=commentaire,
        utilisateur=user if getattr(user, "is_authenticated", False) else None,
    )


def register_diploma_history(diplome, user, action, previous_status="", current_status="", commentaire=""):
    return DiplomeHistory.objects.create(
        diplome=diplome,
        action=action,
        ancien_statut=previous_status,
        nouveau_statut=current_status,
        commentaire=commentaire,
        utilisateur=user if getattr(user, "is_authenticated", False) else None,
    )


# ------------------------------------------------------------ notifications


def get_signature_recipients():
    """Utilisateurs a prevenir lorsqu'un lot part a la signature."""
    return User.objects.filter(
        is_active=True,
        profil__actif=True,
        profil__role__in=[ROLE_DIRECTEUR_GENERAL, ROLE_ADMINISTRATEUR],
    ).distinct()


def get_lot_followers(lot):
    """Agents a prevenir du retour d'un lot : receptionnaire, transmetteur, createur."""
    user_ids = {
        lot.agent_receptionnaire_id,
        lot.transmis_par_id,
        lot.cree_par_id,
    }
    user_ids.discard(None)
    if not user_ids:
        return User.objects.none()
    return User.objects.filter(id__in=user_ids, is_active=True).distinct()


def notify_lot(recipients, titre, message, lot):
    recipients = list(recipients)
    if not recipients:
        return 0

    lot_url = reverse("diplomas:lot_detail", args=[lot.pk])
    Notification.objects.bulk_create(
        [
            Notification(
                utilisateur=user,
                type_notification=Notification.Type.DIPLOME,
                titre=titre,
                message=message,
                url=lot_url,
            )
            for user in recipients
        ]
    )
    return len(recipients)


# -------------------------------------------------------------- import fichier


def normalize_header(value):
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    # Les separateurs (tirets, points, parentheses) deviennent des espaces,
    # puis on ne conserve que les mots : « N° du diplome » -> ["n", "du", "diplome"].
    return "".join(character if character.isalnum() else " " for character in text).split()


def match_column(raw_header):
    normalized = " ".join(normalize_header(raw_header))
    if not normalized:
        return None
    for field, aliases in COLUMN_ALIASES.items():
        if normalized in aliases:
            return field
    for field, aliases in COLUMN_ALIASES.items():
        if any(normalized.startswith(alias) for alias in aliases):
            return field
    return None


def _read_csv_rows(uploaded_file):
    raw_bytes = uploaded_file.read()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Encodage de fichier non reconnu.")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [list(row) for row in reader]


def _read_xlsx_rows(uploaded_file):
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - depend de l'environnement
        raise ValueError(
            "La lecture des fichiers Excel necessite le paquet openpyxl. "
            "Utilisez un fichier CSV ou installez les dependances."
        ) from error

    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = []
        for excel_row in worksheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell).strip() for cell in excel_row])
        return rows
    finally:
        workbook.close()


def read_import_rows(uploaded_file):
    extension = Path(uploaded_file.name).suffix.lower()
    if extension not in SUPPORTED_IMPORT_EXTENSIONS:
        raise ValueError(
            "Format non pris en charge. Fournissez un fichier "
            f"{' ou '.join(SUPPORTED_IMPORT_EXTENSIONS)}."
        )

    if extension == ".csv":
        return _read_csv_rows(uploaded_file)
    return _read_xlsx_rows(uploaded_file)


def parse_diploma_file(uploaded_file, lot):
    """Analyse un fichier importe et retourne des lignes previsualisables.

    Retourne (rows, resume). Aucune ecriture en base n'est effectuee ici :
    l'utilisateur valide l'apercu avant l'enregistrement.
    """
    raw_rows = read_import_rows(uploaded_file)
    if not raw_rows:
        raise ValueError("Le fichier est vide.")

    header_index = None
    column_map = {}
    for index, row in enumerate(raw_rows[:10]):
        candidate = {}
        for position, cell in enumerate(row):
            field = match_column(cell)
            if field and field not in candidate.values():
                candidate[position] = field
        if all(required in candidate.values() for required in REQUIRED_COLUMNS):
            header_index = index
            column_map = candidate
            break

    if header_index is None:
        raise ValueError(
            "Colonnes obligatoires introuvables. Le fichier doit contenir au minimum "
            "une colonne « Nom » et une colonne « Numero »."
        )

    data_rows = raw_rows[header_index + 1 :]
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise ValueError(
            f"Le fichier contient {len(data_rows)} lignes. "
            f"Maximum autorise : {MAX_IMPORT_ROWS} lignes par import."
        )

    existing_numbers = {
        number.strip().lower()
        for number in lot.diplomes.values_list("numero_diplome", flat=True)
    }
    seen_numbers = {}
    rows = []

    for offset, raw_row in enumerate(data_rows):
        line_number = header_index + offset + 2
        values = {field: "" for field in COLUMN_ALIASES}
        for position, field in column_map.items():
            if position < len(raw_row):
                values[field] = str(raw_row[position] or "").strip()

        if not any(values.values()):
            continue

        errors = []
        for field in REQUIRED_COLUMNS:
            if not values[field]:
                label = "Nom du beneficiaire" if field == "nom_beneficiaire" else "Numero du diplome"
                errors.append(f"{label} manquant.")

        for field, max_length in FIELD_MAX_LENGTHS.items():
            if len(values[field]) > max_length:
                errors.append(f"Champ trop long ({field}), maximum {max_length} caracteres.")
                values[field] = values[field][:max_length]

        number_key = values["numero_diplome"].strip().lower()
        if number_key:
            if number_key in existing_numbers:
                errors.append("Ce numero existe deja dans le lot.")
            elif number_key in seen_numbers:
                errors.append(f"Numero en doublon avec la ligne {seen_numbers[number_key]} du fichier.")
            else:
                seen_numbers[number_key] = line_number

        rows.append(
            {
                "line": line_number,
                "data": values,
                "errors": errors,
                "is_valid": not errors,
            }
        )

    if not rows:
        raise ValueError("Aucune ligne exploitable trouvee sous les en-tetes du fichier.")

    resume = {
        "total": len(rows),
        "valides": sum(1 for row in rows if row["is_valid"]),
        "rejetes": sum(1 for row in rows if not row["is_valid"]),
        "colonnes": sorted(set(column_map.values())),
    }
    return rows, resume


def create_diplomas_from_rows(lot, rows, user):
    """Cree les diplomes valides issus d'un apercu d'import."""
    existing_numbers = {
        number.strip().lower()
        for number in lot.diplomes.values_list("numero_diplome", flat=True)
    }
    created = []
    ignored = 0

    for row in rows:
        data = row.get("data", {})
        nom = (data.get("nom_beneficiaire") or "").strip()
        numero = (data.get("numero_diplome") or "").strip()
        if not nom or not numero or numero.lower() in existing_numbers:
            ignored += 1
            continue

        existing_numbers.add(numero.lower())
        created.append(
            Diplome(
                lot=lot,
                nom_beneficiaire=nom[: FIELD_MAX_LENGTHS["nom_beneficiaire"]],
                numero_diplome=numero[: FIELD_MAX_LENGTHS["numero_diplome"]],
                filiere=(data.get("filiere") or "").strip()[: FIELD_MAX_LENGTHS["filiere"]],
                etablissement=(data.get("etablissement") or "").strip()[: FIELD_MAX_LENGTHS["etablissement"]],
                annee_academique=(data.get("annee_academique") or "").strip()[: FIELD_MAX_LENGTHS["annee_academique"]],
                observations=(data.get("observations") or "").strip(),
                statut=Diplome.Status.A_VERIFIER,
            )
        )

    if created:
        Diplome.objects.bulk_create(created)
        register_lot_history(
            lot,
            user,
            f"Import de {len(created)} diplôme(s)",
            lot.statut,
            lot.statut,
            commentaire=f"{ignored} ligne(s) ignorée(s)." if ignored else "",
        )

    return len(created), ignored
